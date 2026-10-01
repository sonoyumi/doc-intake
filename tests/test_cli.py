import pytest
from conftest import EXAMPLES
from openpyxl import load_workbook

from doc_intake import cli
from doc_intake.config import ConfigError, load_settings

KEYS = ("SECRET_KEY", "AI_MODE", "ANTHROPIC_API_KEY", "DATA_DIR", "PORT", "MAX_PAGES", "AI_DAILY_LIMIT")


@pytest.fixture
def env(tmp_path, monkeypatch):
    for key in KEYS:
        monkeypatch.delenv(key, raising=False)
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("SECRET_KEY", "s" * 40)
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setattr(cli.getpass, "getpass", lambda prompt="": "a long enough password")
    return tmp_path


def test_every_config_error_at_once(env, monkeypatch):
    monkeypatch.setenv("SECRET_KEY", "short")
    monkeypatch.setenv("AI_MODE", "claude")
    monkeypatch.setenv("PORT", "eighty")
    monkeypatch.setenv("MAX_PAGES", "500")
    with pytest.raises(ConfigError) as caught:
        load_settings(None)
    text = str(caught.value)
    assert "SECRET_KEY" in text and "ANTHROPIC_API_KEY" in text and "PORT" in text and "MAX_PAGES" in text
    assert cli.main(["stats"]) == 2


def test_users_import_read_export_stats(env, capsys):
    assert cli.main(["add-user", "--email", "anna@studio.it", "--name", "Anna"]) == 0
    assert cli.main(["add-user", "--email", "ANNA@studio.it", "--name", "Again"]) == 1  # e-mail is unique
    assert cli.main(["passwd", "anna@studio.it"]) == 0
    assert cli.main(["passwd", "nobody@studio.it"]) == 1
    assert cli.main(["import", str(EXAMPLES)]) == 0
    assert "Queued: 5 · already known: 0 · skipped: 0" in capsys.readouterr().out
    assert cli.main(["import", str(EXAMPLES / "ddt_ortofrutta.jpg")]) == 0
    assert "Queued: 0 · already known: 1" in capsys.readouterr().out
    assert cli.main(["read", str(EXAMPLES / "fattura_tipografia_alpina.pdf")]) == 0
    assert "All checks passed." in capsys.readouterr().out
    assert cli.main(["stats"]) == 0
    assert "queued: 5" in capsys.readouterr().out
    assert cli.main(["export", "-o", "out.xlsx"]) == 0
    assert "0 approved document(s)" in capsys.readouterr().out
    assert load_workbook(env / "out.xlsx").sheetnames == ["Documents", "Lines", "VAT summary"]
